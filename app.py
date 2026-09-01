import io
from datetime import datetime
import textwrap
import pandas as pd
import plotly.express as px
import streamlit as st
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

try:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import cm
    from reportlab.pdfgen import canvas
except Exception:
    letter = None
    cm = None
    canvas = None

st.set_page_config(
    page_title="Catastro de Vehículos 2027",
    page_icon="🚗",
    layout="wide",
)

ARCHIVO_DEFAULT = "catastro_vehiculos_transformado.xlsx"
HOJA_DEFAULT = "Catastro_final"

CRITERIOS = [
    "Criterio antigüedad >= 8 años",
    "Criterio priorización alta",
    "Criterio gasto mantención > $5.000.000",
    "Criterio kilometraje > 100.000",
]


def normalizar_texto(valor):
    if pd.isna(valor):
        return ""
    return str(valor).strip().lower()


def cumple_criterio(valor):
    texto = normalizar_texto(valor)
    return texto in ["cumple", "si cumple", "sí cumple"]


def convertir_numero(valor):
    if pd.isna(valor):
        return pd.NA

    if isinstance(valor, (int, float)):
        return valor

    texto = str(valor).strip()

    if texto == "" or texto.lower() in [
        "no indica",
        "sin dato",
        "nan",
        "none",
    ]:
        return pd.NA

    texto = texto.replace("$", "").replace(" ", "")

    # Formato chileno con miles y decimales, por ejemplo 1.234.567,89
    if "." in texto and "," in texto:
        texto = texto.replace(".", "").replace(",", ".")

    # Solo coma decimal, por ejemplo 2018,0
    elif "," in texto:
        texto = texto.replace(",", ".")

    # Si tiene varios puntos, son separadores de miles
    elif texto.count(".") > 1:
        texto = texto.replace(".", "")

    return pd.to_numeric(texto, errors="coerce")


def formato_miles(valor):
    numero = pd.to_numeric(valor, errors="coerce")
    if pd.isna(numero):
        return "No indica"
    return f"{numero:,.0f}".replace(",", ".")


def formato_pesos(valor):
    numero = pd.to_numeric(valor, errors="coerce")
    if pd.isna(numero):
        return "No indica"
    return "$" + f"{numero:,.0f}".replace(",", ".")


@st.cache_data(show_spinner=False)
def leer_excel(archivo):
    """Lee la hoja Catastro_final. Si no existe, usa la primera hoja disponible."""
    xls = pd.ExcelFile(archivo, engine="openpyxl")
    hoja = HOJA_DEFAULT if HOJA_DEFAULT in xls.sheet_names else xls.sheet_names[0]
    data = pd.read_excel(archivo, sheet_name=hoja, engine="openpyxl")
    data.columns = [str(c).strip() for c in data.columns]
    return data, hoja

@st.cache_data(show_spinner=False)
def leer_dotacion_maxima(archivo):
    nombre_hoja = "Dotación_máxima"

    dotacion = pd.read_excel(
        archivo,
        sheet_name=nombre_hoja,
        engine="openpyxl"
    )

    dotacion.columns = [
        str(columna).strip()
        for columna in dotacion.columns
    ]

    # Eliminar filas completamente vacías
    dotacion = dotacion.dropna(how="all")

    # Normalizar texto del Servicio
    dotacion["Servicio"] = (
        dotacion["Servicio"]
        .astype(str)
        .str.strip()
    )

    # Convertir dotación a valor numérico
    dotacion["Dotacion máxima autorizada"] = pd.to_numeric(
        dotacion["Dotacion máxima autorizada"],
        errors="coerce"
    )

    # Excluir filas de totales
    dotacion = dotacion[
        ~dotacion["Servicio"].str.lower().str.startswith("total")
    ]

    # Excluir filas sin dotación válida
    dotacion = dotacion.dropna(
        subset=["Dotacion máxima autorizada"]
    )

    return dotacion

def preparar_datos(df):
    df = df.copy()

   # Corregir año vehículo para eliminar decimales
    if "Año Vehículo" in df.columns:
      df["Año Vehículo"] = (
        pd.to_numeric(df["Año Vehículo"], errors="coerce")
        .astype("Int64")
    )

    if "Tipo vehículo" in df.columns:
        df["Tipo vehículo"] = (
            df["Tipo vehículo"]
            .fillna("")
            .astype(str)
            .str.strip()
            .str.title()
            .replace({
                "Camiioneta": "Camioneta",
                "Mini Bus": "Minibus",
            })
        )
        df.loc[df["Tipo vehículo"].isin(["", "Nan", "None"]), "Tipo vehículo"] = "No indica"

    for col in ["Kilometraje acumulado", "Gasto en mantención acumulada", "Año Vehículo"]:
        if col in df.columns:
            df[col + " num"] = df[col].apply(convertir_numero)

    if "Año Vehículo num" in df.columns:
        anio_valido = df["Año Vehículo num"].where(
            df["Año Vehículo num"].between(1900, 2027)
        )

        df["Antigüedad 2027"] = 2027 - anio_valido

    # Recalcular criterios con base en los datos de cada vehículo
    if "Antigüedad 2027" in df.columns:
        df["Criterio antigüedad >= 8 años"] = df["Antigüedad 2027"].apply(
            lambda valor: (
                "cumple"
                if pd.notna(valor) and valor >= 8
                else "no cumple"
            )
        )

    if "Priorización" in df.columns:
        df["Criterio priorización alta"] = df["Priorización"].apply(
            lambda valor: (
                "cumple"
                if normalizar_texto(valor) == "alta"
                else "no cumple"
            )
        )

    if "Gasto en mantención acumulada num" in df.columns:
        df["Criterio gasto mantención > $5.000.000"] = df[
            "Gasto en mantención acumulada num"
        ].apply(
            lambda valor: (
                "si cumple"
                if pd.notna(valor) and valor > 5_000_000
                else "no cumple"
            )
        )

    if "Kilometraje acumulado num" in df.columns:
        df["Criterio kilometraje > 100.000"] = df[
            "Kilometraje acumulado num"
        ].apply(
            lambda valor: (
                "si cumple"
                if pd.notna(valor) and valor > 100_000
                else "no cumple"
            )
        )

    for criterio in CRITERIOS:
        if criterio not in df.columns:
            df[criterio] = "no cumple"

    df["Cantidad criterios cumplidos"] = df[CRITERIOS].apply(
        lambda fila: sum(cumple_criterio(v) for v in fila), axis=1
    )

    df["Cumple 3 o más criterios"] = df["Cantidad criterios cumplidos"].apply(
        lambda x: "cumple" if x >= 3 else "no cumple"
    )

    df["Cumple 4 criterios"] = df["Cantidad criterios cumplidos"].apply(
        lambda x: "cumple" if x == 4 else "no cumple"
    )

    df["Cumple antigüedad y priorización alta"] = df.apply(
        lambda r: "cumple"
        if cumple_criterio(r.get("Criterio antigüedad >= 8 años"))
        and cumple_criterio(r.get("Criterio priorización alta"))
        else "no cumple",
        axis=1,
    )
    return df

def aplicar_multiselect(df, columna, etiqueta):
    if columna not in df.columns:
        return df

    if columna == "Año Vehículo":
        opciones = sorted(
            int(x) for x in df[columna].dropna().unique()
        )

        seleccion = st.sidebar.multiselect(
            etiqueta,
            opciones
        )

        if seleccion:
            return df[
                pd.to_numeric(df[columna], errors="coerce")
                .astype("Int64")
                .isin(seleccion)
            ]

    else:
        opciones = sorted(
            str(x) for x in df[columna].dropna().unique()
        )

        seleccion = st.sidebar.multiselect(
            etiqueta,
            opciones
        )

        if seleccion:
            return df[
                df[columna].astype(str).isin(seleccion)
            ]

    return df

def dataframe_a_excel(df):
    salida = io.BytesIO()
    exportar = df.copy()
    cols_ocultar = [c for c in exportar.columns if c.endswith(" num")]
    exportar = exportar.drop(columns=cols_ocultar, errors="ignore")

    with pd.ExcelWriter(salida, engine="openpyxl") as writer:
        exportar.to_excel(writer, index=False, sheet_name="Reporte_filtrado")
        ws = writer.book["Reporte_filtrado"]

        # Estilo encabezado
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        encabezados = {cell.value: cell.column for cell in ws[1]}
        for col_nombre in ["Kilometraje acumulado", "Año Vehículo", "Año de Compra"]:
            if col_nombre in encabezados:
                col_idx = encabezados[col_nombre]
                for row in range(2, ws.max_row + 1):
                    ws.cell(row=row, column=col_idx).number_format = '#,##0'

        if "Gasto en mantención acumulada" in encabezados:
            col_idx = encabezados["Gasto en mantención acumulada"]
            for row in range(2, ws.max_row + 1):
                ws.cell(row=row, column=col_idx).number_format = '$#,##0'

        for col_idx, column_cells in enumerate(ws.columns, start=1):
            max_len = 0
            for cell in column_cells:
                max_len = max(max_len, len(str(cell.value)) if cell.value is not None else 0)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 12), 45)

    salida.seek(0)
    return salida


def generar_minuta(
    df,
    filtros_aplicados,
    resumen_dotacion=None
):   
    total = len(df)
    servicios = df["Servicio"].nunique()

    # ---------------------------------------------------------
    # Funciones auxiliares utilizadas solamente en la minuta
    # ---------------------------------------------------------

    def obtener_texto(valor, texto_default="No indica"):
        if pd.isna(valor):
            return texto_default

        texto = str(valor).strip()

        if texto.lower() in ["", "nan", "none"]:
            return texto_default

        return texto

    def obtener_anio(valor):
        if pd.isna(valor):
            return "No indica"

        try:
            return str(int(float(valor)))
        except (TypeError, ValueError):
            return obtener_texto(valor)

    def obtener_resultado_criterio(valor):
        return "Cumple" if cumple_criterio(valor) else "No cumple"

    def obtener_servicios(dataframe):
        if "Servicio" not in dataframe.columns:
            return []

        return (
            dataframe["Servicio"]
            .dropna()
            .astype(str)
            .str.strip()
            .replace("", pd.NA)
            .dropna()
            .unique()
            .tolist()
        )

    servicios_unicos = obtener_servicios(df)

    if len(servicios_unicos) == 1:
        nombre_servicio = servicios_unicos[0]
    elif len(servicios_unicos) > 1:
        nombre_servicio = f"{len(servicios_unicos)} servicios institucionales"
    else:
        nombre_servicio = "servicio no identificado"

    lineas = []

    # =========================================================
    # MINUTA SIN REGISTROS
    # =========================================================

    if total == 0:
        lineas.append("MINUTA EJECUTIVA")
        lineas.append("Sección de Planificación y Presupuesto")
        lineas.append("Departamento de Compras y SS.GG. - DIVAD")
        lineas.append("")
        lineas.append(
            "ANÁLISIS DE SOLICITUDES DE REEMPLAZO "
            "DE VEHÍCULOS INSTITUCIONALES"
        )
        lineas.append("")
        lineas.append(
            f"Fecha de emisión: "
            f"{datetime.now().strftime('%d-%m-%Y')}"
        )
        lineas.append("")
        lineas.append("No existen vehículos asociados a los filtros aplicados.")

        return "\n".join(lineas)

    # =========================================================
    # MINUTA INDIVIDUAL
    # =========================================================

    if total == 1:
        fila = df.iloc[0]

        servicio = obtener_texto(
            fila.get("Servicio", pd.NA)
        )

        patente = obtener_texto(
            fila.get("I.R.N.V.M.", pd.NA)
        )

        tipo_vehiculo = obtener_texto(
            fila.get("Tipo vehículo", pd.NA)
        )

        anio_vehiculo = obtener_anio(
            fila.get("Año Vehículo", pd.NA)
        )

        kilometraje = fila.get(
            "Kilometraje acumulado num",
            fila.get("Kilometraje acumulado", pd.NA),
        )

        gasto_mantencion = fila.get(
            "Gasto en mantención acumulada num",
            fila.get("Gasto en mantención acumulada", pd.NA),
        )

        estado = obtener_texto(
            fila.get("Estado", pd.NA)
        )

        priorizacion = obtener_texto(
            fila.get("Priorización", pd.NA)
        )

        cantidad_criterios = int(
            fila.get("Cantidad criterios cumplidos", 0)
        )

        criterio_antiguedad = obtener_resultado_criterio(
            fila.get("Criterio antigüedad >= 8 años", "")
        )

        criterio_priorizacion = obtener_resultado_criterio(
            fila.get("Criterio priorización alta", "")
        )

        criterio_gasto = obtener_resultado_criterio(
            fila.get(
                "Criterio gasto mantención > $5.000.000",
                "",
            )
        )

        criterio_kilometraje = obtener_resultado_criterio(
            fila.get(
                "Criterio kilometraje > 100.000",
                "",
            )
        )

        tipo_falla = obtener_texto(
            fila.get("Tipo de falla", pd.NA)
        )

        justificacion = obtener_texto(
            fila.get("Justificación (descripción)", pd.NA)
        )
        lineas.append("MINUTA EJECUTIVA")
        lineas.append("Sección de Planificación y Presupuesto")
        lineas.append("Departamento de Compras y SS.GG. - DIVAD")
        lineas.append("")
        lineas.append(
            "SOLICITUD DE REEMPLAZO DE VEHÍCULO INSTITUCIONAL"
        )
        lineas.append("")
        lineas.append(
            f"Fecha de emisión: "
            f"{datetime.now().strftime('%d-%m-%Y')}"
        )
        lineas.append(f"Servicio: {servicio}")
        lineas.append(f"I.R.N.V.M.: {patente}")

        lineas.append("")
        lineas.append("1. OBJETIVO")
        lineas.append("")
        lineas.append(
            "Analizar y fundamentar técnicamente la solicitud "
            "de reemplazo del vehículo institucional seleccionado, "
            "considerando los criterios establecidos para el proceso "
            "de renovación vehicular 2027."
        )

        lineas.append("")
        lineas.append("2. ANTECEDENTES DEL VEHÍCULO")
        lineas.append("")
        lineas.append(f"- Servicio: {servicio}")
        lineas.append(f"- I.R.N.V.M.: {patente}")
        lineas.append(f"- Tipo de vehículo: {tipo_vehiculo}")
        lineas.append(f"- Año del vehículo: {anio_vehiculo}")
        lineas.append(
            f"- Kilometraje acumulado: "
            f"{formato_miles(kilometraje)} km"
        )
        lineas.append(
            f"- Gasto en mantención acumulada: "
            f"{formato_pesos(gasto_mantencion)}"
        )
        lineas.append(f"- Estado: {estado}")
        lineas.append(f"- Priorización: {priorizacion}")

        lineas.append("")
        lineas.append("3. EVALUACIÓN DE CRITERIOS")
        lineas.append("")
        lineas.append(
            f"- Antigüedad igual o superior a ocho años: "
            f"{criterio_antiguedad}"
        )
        lineas.append(
            f"- Priorización alta: {criterio_priorizacion}"
        )
        lineas.append(
            f"- Gasto de mantención superior a $5.000.000: "
            f"{criterio_gasto}"
        )
        lineas.append(
            f"- Kilometraje superior a 100.000 km: "
            f"{criterio_kilometraje}"
        )
        lineas.append("")
        lineas.append(
            f"Cantidad de criterios cumplidos: "
            f"{cantidad_criterios} de 4."
        )

        lineas.append("")
        lineas.append("4. ANTECEDENTES TÉCNICOS RELEVANTES")
        lineas.append("")
        lineas.append(f"- Tipo de falla: {tipo_falla}")
        lineas.append(f"- Justificación técnica: {justificacion}")

        lineas.append("")
        lineas.append("5. CONCLUSIÓN TÉCNICA")
        lineas.append("")

        if cantidad_criterios == 4:
            lineas.append(
                "En atención a los antecedentes expuestos, se concluye "
                "que el vehículo analizado cumple con los cuatro criterios "
                "de evaluación establecidos para el proceso de renovación "
                "vehicular. Considerando su antigüedad, priorización, "
                "kilometraje acumulado y gastos de mantención, se estima "
                "técnicamente procedente evaluar favorablemente su "
                "reemplazo, con el objeto de mantener la continuidad "
                "operativa del servicio en condiciones adecuadas de "
                "seguridad, confiabilidad y eficiencia."
            )

        elif cantidad_criterios == 3:
            lineas.append(
                "En atención a los antecedentes expuestos, se observa "
                "que el vehículo analizado cumple con tres de los cuatro "
                "criterios establecidos para el proceso de renovación "
                "vehicular. En consecuencia, se estima que su reemplazo "
                "debe ser evaluado considerando adicionalmente su "
                "condición mecánica, continuidad operativa, disponibilidad "
                "presupuestaria y antecedentes de seguridad."
            )

        else:
            lineas.append(
                "En atención a los antecedentes expuestos, se observa "
                f"que el vehículo analizado cumple con "
                f"{cantidad_criterios} de los cuatro criterios establecidos. "
                "Por lo anterior, no resulta procedente recomendar su "
                "reemplazo únicamente con base en los criterios "
                "cuantitativos. No obstante, la solicitud deberá evaluarse "
                "considerando eventuales fallas graves, riesgos para la "
                "seguridad, falta de repuestos, inactividad prolongada "
                "u otros antecedentes técnicos relevantes."
            )

        lineas.append("")
        lineas.append("6. RECOMENDACIÓN")
        lineas.append("")

        if cantidad_criterios == 4:
            lineas.append(
                "Se recomienda incorporar el vehículo al proceso de "
                "renovación vehicular 2027 y continuar la tramitación "
                "administrativa correspondiente para su evaluación "
                "presupuestaria."
            )

        elif cantidad_criterios == 3:
            lineas.append(
                "Se recomienda complementar los antecedentes disponibles "
                "con una evaluación técnica actualizada y determinar la "
                "pertinencia de incorporar el vehículo al proceso de "
                "renovación vehicular 2027."
            )

        else:
            lineas.append(
                "Se recomienda mantener el vehículo en evaluación y "
                "solicitar antecedentes técnicos adicionales cuando "
                "existan fallas, riesgos operativos o condiciones que "
                "puedan justificar su reposición."
            )

        return "\n".join(lineas)

    # =========================================================
    # MINUTA CONSOLIDADA PARA VARIOS VEHÍCULOS
    # =========================================================

    priorizacion_alta = 0

    if "Priorización" in df.columns:
        priorizacion_alta = (
            df["Priorización"]
            .astype(str)
            .str.strip()
            .str.lower()
            .eq("alta")
            .sum()
        )

    if "Cantidad criterios cumplidos" in df.columns:
        cuatro_criterios = (
            df["Cantidad criterios cumplidos"].eq(4).sum()
        )

        tres_criterios = (
            df["Cantidad criterios cumplidos"].eq(3).sum()
        )

        dos_o_menos = (
            df["Cantidad criterios cumplidos"].le(2).sum()
        )
    else:
        cuatro_criterios = 0
        tres_criterios = 0
        dos_o_menos = total

    gasto_total = df.get(
        "Gasto en mantención acumulada num",
        pd.Series(dtype=float),
    ).sum(skipna=True)

    kilometraje_promedio = df.get(
        "Kilometraje acumulado num",
        pd.Series(dtype=float),
    ).mean(skipna=True)

    lineas.append("MINUTA EJECUTIVA")
    lineas.append("Sección de Planificación y Presupuesto")
    lineas.append("Departamento de Compras y SS.GG. - DIVAD")
    lineas.append("")
    lineas.append(
          "ANÁLISIS CONSOLIDADO DE SOLICITUDES DE REEMPLAZO "
          "DE VEHÍCULOS INSTITUCIONALES"
    )
    lineas.append("")
    lineas.append(
        f"Fecha de emisión: {datetime.now().strftime('%d-%m-%Y')}"
    )
    lineas.append(f"Universo analizado: {total} vehículos")

    if len(servicios_unicos) == 1:
        lineas.append(f"Servicio: {servicios_unicos[0]}")
    else:
        lineas.append(
            f"Servicios incluidos: {len(servicios_unicos)}"
        )

    lineas.append("")
    lineas.append("1. OBJETIVO")
    lineas.append("")
    lineas.append(
        "Analizar y fundamentar técnicamente las solicitudes de "
        "reemplazo de los vehículos institucionales seleccionados, "
        "considerando su antigüedad, nivel de priorización, gasto "
        "acumulado en mantención y kilometraje."
    )

    lineas.append("")
    lineas.append("2. RESUMEN EJECUTIVO")
    lineas.append("")
    lineas.append(
        f"- Total de vehículos analizados: {total}"
    )
    lineas.append(
        f"- Vehículos con priorización alta: {priorizacion_alta}"
    )
    lineas.append(
        f"- Vehículos que cumplen cuatro criterios: "
        f"{cuatro_criterios}"
    )
    lineas.append(
        f"- Vehículos que cumplen tres criterios: "
        f"{tres_criterios}"
    )
    lineas.append(
        f"- Vehículos que cumplen dos o menos criterios: "
        f"{dos_o_menos}"
    )

    if pd.notna(kilometraje_promedio):
        lineas.append(
            f"- Kilometraje promedio: "
            f"{formato_miles(kilometraje_promedio)} km"
        )
    else:
        lineas.append(
            "- Kilometraje promedio: No disponible"
        )

    lineas.append(
        f"- Gasto total de mantención acumulada: "
        f"{formato_pesos(gasto_total)}"
    )

    lineas.append("")
    lineas.append("3. DETALLE DE VEHÍCULOS ANALIZADOS")
    lineas.append("")

    limite_pdf = 50

    for _, fila in df.head(limite_pdf).iterrows():
        servicio = obtener_texto(
            fila.get("Servicio", pd.NA)
        )

        patente = obtener_texto(
            fila.get("I.R.N.V.M.", pd.NA)
        )

        tipo_vehiculo = obtener_texto(
            fila.get("Tipo vehículo", pd.NA)
        )

        anio_vehiculo = obtener_anio(
            fila.get("Año Vehículo", pd.NA)
        )

        kilometraje = fila.get(
            "Kilometraje acumulado num",
            fila.get("Kilometraje acumulado", pd.NA),
        )

        gasto = fila.get(
            "Gasto en mantención acumulada num",
            fila.get("Gasto en mantención acumulada", pd.NA),
        )

        cantidad_criterios = int(
            fila.get("Cantidad criterios cumplidos", 0)
        )

        if cantidad_criterios == 4:
            resultado = "Reemplazo recomendado"
        elif cantidad_criterios == 3:
            resultado = "Evaluar reemplazo"
        else:
            resultado = "Revisar antecedentes"

        lineas.append(
            f"- Servicio: {servicio} | "
            f"I.R.N.V.M.: {patente} | "
            f"Tipo: {tipo_vehiculo} | "
            f"Año: {anio_vehiculo} | "
            f"Kilometraje: {formato_miles(kilometraje)} km | "
            f"Gasto mantención: {formato_pesos(gasto)} | "
            f"Criterios: {cantidad_criterios} de 4 | "
            f"Resultado: {resultado}"
        )

    if total > limite_pdf:
        lineas.append("")
        lineas.append(
            f"Se omiten {total - limite_pdf} registros adicionales "
            "por extensión de la minuta. La nómina completa se encuentra "
            "disponible en el archivo Excel descargable."
        )

    lineas.append("")
    lineas.append("4. EVALUACIÓN CONSOLIDADA")
    lineas.append("")
    lineas.append("4.1 Vehículos que cumplen los cuatro criterios")
    lineas.append("")
    lineas.append(
        "Los vehículos que cumplen los cuatro criterios presentan "
        "condiciones objetivas que justifican su incorporación "
        "prioritaria al proceso de renovación vehicular. En estos casos "
        "concurren simultáneamente la antigüedad igual o superior a ocho "
        "años, la priorización alta, un gasto acumulado en mantención "
        "superior a $5.000.000 y un kilometraje acumulado superior a "
        "100.000 kilómetros."
    )

    lineas.append("")
    lineas.append("4.2 Vehículos que cumplen tres criterios")
    lineas.append("")
    lineas.append(
        "Los vehículos que cumplen tres criterios presentan antecedentes "
        "relevantes para considerar su eventual reemplazo. La decisión "
        "deberá complementarse con una revisión de su condición mecánica, "
        "continuidad operativa, disponibilidad presupuestaria y "
        "justificación técnica."
    )

    lineas.append("")
    lineas.append(
        "4.3 Vehículos que cumplen dos o menos criterios"
    )
    lineas.append("")
    lineas.append(
        "Los vehículos que cumplen dos o menos criterios no presentan, "
        "únicamente sobre la base de los criterios cuantitativos, "
        "antecedentes suficientes para recomendar directamente su "
        "reemplazo. No obstante, estos casos deberán evaluarse "
        "individualmente cuando existan fallas mecánicas graves, riesgos "
        "para la seguridad, falta de repuestos, inactividad prolongada "
        "u otros antecedentes técnicos relevantes."
    )

    lineas.append("")
    lineas.append("5. CONCLUSIÓN")
    lineas.append("")
    lineas.append(
        "En atención al análisis efectuado, se concluye que el universo "
        "seleccionado presenta distintos niveles de cumplimiento de los "
        "criterios establecidos para el proceso de renovación vehicular."
    )
    lineas.append("")
    lineas.append(
        "Respecto de los vehículos que cumplen los cuatro criterios, "
        "se estima técnicamente procedente priorizar su reemplazo, con "
        "el propósito de mantener la continuidad operativa de los "
        "servicios en condiciones adecuadas de seguridad, confiabilidad "
        "y eficiencia."
    )
    lineas.append("")
    lineas.append(
        "En cuanto a los vehículos que no cumplen la totalidad de los "
        "criterios, se recomienda evaluar la pertinencia de su reemplazo "
        "caso a caso, considerando los antecedentes técnicos, operativos, "
        "presupuestarios y de seguridad disponibles."
    )

        # ----------------------------------------------------------
    # Información de dotación máxima
    # ----------------------------------------------------------
    lineas.append("")
    lineas.append("6. DOTACIÓN MÁXIMA VEHICULAR")
    lineas.append("")

    if (
        resumen_dotacion is not None
        and not resumen_dotacion.empty
    ):
        dotacion_actual = resumen_dotacion[
            "Vehículos actuales"
        ].sum()

        dotacion_autorizada = resumen_dotacion[
            "Dotacion máxima autorizada"
        ].sum()

        cupos_disponibles = (
            dotacion_autorizada
            - dotacion_actual
        )

        ocupacion_dotacion = (
            dotacion_actual / dotacion_autorizada * 100
            if dotacion_autorizada > 0
            else 0
        )

        lineas.append(
            f"- Vehículos actuales: "
            f"{dotacion_actual:,.0f}".replace(",", ".")
        )

        lineas.append(
            f"- Dotación máxima autorizada: "
            f"{dotacion_autorizada:,.0f}".replace(",", ".")
        )

        lineas.append(
            f"- Cupos disponibles: "
            f"{cupos_disponibles:,.0f}".replace(",", ".")
        )

        lineas.append(
            f"- Porcentaje de ocupación: "
            f"{ocupacion_dotacion:.1f}%"
        )

        lineas.append("")
        lineas.append("Detalle por Servicio:")

        for _, fila_dotacion in resumen_dotacion.iterrows():
            servicio_dotacion = fila_dotacion.get(
                "Servicio",
                "No identificado"
            )

            actuales = fila_dotacion.get(
                "Vehículos actuales",
                0
            )

            autorizada = fila_dotacion.get(
                "Dotacion máxima autorizada",
                0
            )

            cupos = fila_dotacion.get(
                "Cupos disponibles",
                autorizada - actuales
            )

            ocupacion = fila_dotacion.get(
                "Porcentaje de ocupación",
                0
            )

            situacion = fila_dotacion.get(
                "Situación",
                "No indica"
            )

            if pd.isna(ocupacion):
                ocupacion = 0

            lineas.append(
                f"- {servicio_dotacion}: "
                f"{actuales:.0f} vehículos actuales; "
                f"dotación máxima autorizada {autorizada:.0f}; "
                f"{cupos:.0f} cupos disponibles; "
                f"ocupación {ocupacion:.1f}%; "
                f"situación: {situacion}."
            )

    else:
        lineas.append(
            "- No existe información de dotación máxima "
            "para los servicios seleccionados."
        )

    lineas.append("")
    lineas.append("7. RECOMENDACIONES")
    lineas.append("")
    lineas.append(
        "1. Priorizar los vehículos que cumplen los cuatro criterios."
    )
    lineas.append(
        "2. Evaluar individualmente los vehículos que cumplen tres "
        "criterios."
    )
    lineas.append(
        "3. Solicitar antecedentes técnicos adicionales para los "
        "vehículos con fallas graves que cumplen dos o menos criterios."
    )
    lineas.append(
        "4. Revisar la disponibilidad presupuestaria para determinar "
        "la programación anual de reposiciones."
    )
    lineas.append(
        "5. Mantener trazabilidad de las decisiones adoptadas respecto "
        "de cada vehículo."
    )
    # =========================================================
    # VEHÍCULOS PRIORITARIOS
    # =========================================================

    lineas.append("")
    lineas.append("8. VEHÍCULOS PRIORITARIOS")
    lineas.append("")

    if "Cantidad criterios cumplidos" in df.columns:

        vehiculos_prioritarios = df[
            df["Cantidad criterios cumplidos"] == 4
        ].copy()

        if not vehiculos_prioritarios.empty:

            # Crear columnas auxiliares para ordenar correctamente
            vehiculos_prioritarios["_prioridad_orden"] = (
                vehiculos_prioritarios["Priorización"]
                .astype(str)
                .str.strip()
                .str.lower()
                .map({
                    "alta": 1,
                    "media": 2,
                    "baja": 3,
                })
                .fillna(4)
            )

            vehiculos_prioritarios["_gasto_orden"] = pd.to_numeric(
                vehiculos_prioritarios.get(
                    "Gasto en mantención acumulada num",
                    pd.Series(
                        0,
                        index=vehiculos_prioritarios.index,
                    ),
                ),
                errors="coerce",
            ).fillna(0)

            vehiculos_prioritarios["_km_orden"] = pd.to_numeric(
                vehiculos_prioritarios.get(
                    "Kilometraje acumulado num",
                    pd.Series(
                        0,
                        index=vehiculos_prioritarios.index,
                    ),
                ),
                errors="coerce",
            ).fillna(0)

            vehiculos_prioritarios["_anio_orden"] = pd.to_numeric(
                vehiculos_prioritarios.get(
                    "Año Vehículo num",
                    pd.Series(
                        9999,
                        index=vehiculos_prioritarios.index,
                    ),
                ),
                errors="coerce",
            ).fillna(9999)

            vehiculos_prioritarios = vehiculos_prioritarios.sort_values(
                by=[
                    "_prioridad_orden",
                    "_gasto_orden",
                    "_km_orden",
                    "_anio_orden",
                ],
                ascending=[
                    True,
                    False,
                    False,
                    True,
                ],
            )

            lineas.append(
                "De acuerdo con los criterios aplicados, se identifican "
                "los siguientes vehículos como prioritarios para el "
                "proceso de renovación vehicular:"
            )
            lineas.append("")

            limite_prioritarios = 20

            for numero, (_, fila) in enumerate(
                vehiculos_prioritarios.head(
                    limite_prioritarios
                ).iterrows(),
                start=1,
            ):

                servicio = obtener_texto(
                    fila.get("Servicio", pd.NA)
                )

                patente = obtener_texto(
                    fila.get("I.R.N.V.M.", pd.NA)
                )

                tipo_vehiculo = obtener_texto(
                    fila.get("Tipo vehículo", pd.NA)
                )

                anio_vehiculo = obtener_anio(
                    fila.get("Año Vehículo", pd.NA)
                )

                priorizacion = obtener_texto(
                    fila.get("Priorización", pd.NA)
                )

                kilometraje = fila.get(
                    "Kilometraje acumulado num",
                    fila.get(
                        "Kilometraje acumulado",
                        pd.NA,
                    ),
                )

                gasto = fila.get(
                    "Gasto en mantención acumulada num",
                    fila.get(
                        "Gasto en mantención acumulada",
                        pd.NA,
                    ),
                )

                lineas.append(
                    f"{numero}. {patente} - {servicio}"
                )
                lineas.append(
                    f"   Tipo de vehículo: {tipo_vehiculo}"
                )
                lineas.append(
                    f"   Año del vehículo: {anio_vehiculo}"
                )
                lineas.append(
                    f"   Priorización: {priorizacion}"
                )
                lineas.append(
                    f"   Kilometraje acumulado: "
                    f"{formato_miles(kilometraje)} km"
                )
                lineas.append(
                    f"   Gasto en mantención acumulada: "
                    f"{formato_pesos(gasto)}"
                )
                lineas.append(
                    "   Criterios cumplidos: 4 de 4"
                )
                lineas.append("")

            if len(vehiculos_prioritarios) > limite_prioritarios:
                lineas.append(
                    f"Se omiten "
                    f"{len(vehiculos_prioritarios) - limite_prioritarios} "
                    "vehículos prioritarios adicionales por extensión "
                    "de la minuta. La nómina completa se encuentra "
                    "disponible en el archivo Excel descargable."
                )

        else:
            lineas.append(
                "En el universo seleccionado no se identifican "
                "vehículos que cumplan los cuatro criterios de "
                "evaluación establecidos."
            )

    else:
        lineas.append(
            "No se dispone de información suficiente para identificar "
            "vehículos prioritarios."
        )

    return "\n".join(lineas)


def minuta_a_pdf_bytes(texto):
    if canvas is None:
        raise RuntimeError("No está disponible la librería reportlab. Agregue reportlab a requirements.txt.")

    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=letter)
    ancho, alto = letter
    margen_x = 2 * cm
    margen_y = 2 * cm
    y = alto - margen_y
    ancho_linea = 95

    c.setFont("Helvetica", 10)

    for parrafo in texto.split("\n"):
        lineas = textwrap.wrap(parrafo, width=ancho_linea) if parrafo else [""]
        for linea in lineas:
            if y < margen_y:
                c.showPage()
                y = alto - margen_y
                c.setFont("Helvetica", 10)
            c.drawString(margen_x, y, linea)
            y -= 0.45 * cm

    c.save()
    buffer.seek(0)
    return buffer


st.title("🚗 Catastro de Vehículos")
st.caption("Aplicación para análisis, filtros, gráficos y generación de minuta del catastro de vehículos.")

with st.sidebar:
    st.header("Carga de datos")
    archivo_subido = st.file_uploader("Cargar archivo Excel", type=["xlsx"])

try:
    archivo = (
        archivo_subido
        if archivo_subido is not None
        else ARCHIVO_DEFAULT
    )

    df_original, hoja_usada = leer_excel(archivo)
    df_dotacion = leer_dotacion_maxima(archivo)

except Exception as e:
    st.error(
        "No fue posible cargar el archivo Excel o la hoja "
        "'Dotación_máxima'. Revise los nombres de las hojas "
        "y columnas."
    )
    st.exception(e)
    st.stop()

except Exception as e:
    st.error("No fue posible cargar el archivo Excel. Verifique que el archivo esté en la misma carpeta de app.py o cárguelo desde la barra lateral.")
    st.exception(e)
    st.stop()

st.info(f"Hoja utilizada: {hoja_usada} | Registros cargados: {len(df_original):,}".replace(",", "."))

df = preparar_datos(df_original)

# Calcular cantidad actual de vehículos por Servicio
vehiculos_actuales = (
    df.groupby("Servicio")
    .size()
    .reset_index(name="Vehículos actuales")
)

# Relacionar el catastro con la dotación máxima
resumen_dotacion = df_dotacion.merge(
    vehiculos_actuales,
    on="Servicio",
    how="left"
)

resumen_dotacion["Vehículos actuales"] = (
    resumen_dotacion["Vehículos actuales"]
    .fillna(0)
    .astype(int)
)

resumen_dotacion["Dotacion máxima autorizada"] = (
    resumen_dotacion["Dotacion máxima autorizada"]
    .astype(int)
)

# Calcular cupos disponibles
resumen_dotacion["Cupos disponibles"] = (
    resumen_dotacion["Dotacion máxima autorizada"]
    - resumen_dotacion["Vehículos actuales"]
)

# Calcular porcentaje de ocupación
resumen_dotacion["Porcentaje de ocupación"] = (
    resumen_dotacion["Vehículos actuales"]
    / resumen_dotacion["Dotacion máxima autorizada"]
    * 100
).round(1)

# Clasificar situación de cada Servicio
resumen_dotacion["Situación"] = resumen_dotacion[
    "Cupos disponibles"
].apply(
    lambda valor: (
        "Sobre dotación"
        if valor < 0
        else "Dotación completa"
        if valor == 0
        else "Con cupos disponibles"
    )
)


st.sidebar.header("Filtros")
df_filtrado = df.copy()
filtros_aplicados = []

for columna, etiqueta in [
    ("Servicio", "Servicio"),
    ("Año de Compra", "Año de Compra"),
    ("I.R.N.V.M.", "I.R.N.V.M."),
    ("Año Vehículo", "Año Vehículo"),
    ("Tipo vehículo", "Tipo de vehículo"),
    ("Estado", "Estado"),
    ("Priorización", "Priorización"),
    ("Criterio antigüedad >= 8 años", "Criterio antigüedad"),
    ("Criterio priorización alta", "Criterio priorización alta"),
    ("Criterio gasto mantención > $5.000.000", "Criterio gasto mantención"),
    ("Criterio kilometraje > 100.000", "Criterio kilometraje"),
]:
    antes = len(df_filtrado)
    df_filtrado = aplicar_multiselect(df_filtrado, columna, etiqueta)
    if len(df_filtrado) != antes:
        filtros_aplicados.append(f"{etiqueta}: filtro aplicado")

# ==========================================================
# DOTACIÓN MÁXIMA DINÁMICA SEGÚN LOS FILTROS
# ==========================================================

st.subheader("Dotación máxima por Servicio")

# Contar los vehículos que permanecen después de aplicar filtros
vehiculos_actuales_filtrados = (
    df_filtrado.groupby("Servicio")
    .size()
    .reset_index(name="Vehículos actuales")
)

# Obtener los servicios presentes en el resultado filtrado
servicios_filtrados = (
    df_filtrado["Servicio"]
    .dropna()
    .astype(str)
    .str.strip()
    .unique()
)

# Considerar solamente la dotación de los servicios filtrados
dotacion_filtrada = df_dotacion[
    df_dotacion["Servicio"].isin(servicios_filtrados)
].copy()

# Relacionar dotación máxima con vehículos filtrados
resumen_dotacion_filtrado = dotacion_filtrada.merge(
    vehiculos_actuales_filtrados,
    on="Servicio",
    how="left"
)

# Completar servicios que no tengan vehículos en el resultado
resumen_dotacion_filtrado["Vehículos actuales"] = (
    resumen_dotacion_filtrado["Vehículos actuales"]
    .fillna(0)
    .astype(int)
)

resumen_dotacion_filtrado["Dotacion máxima autorizada"] = (
    pd.to_numeric(
        resumen_dotacion_filtrado[
            "Dotacion máxima autorizada"
        ],
        errors="coerce"
    )
    .fillna(0)
    .astype(int)
)

# Calcular cupos disponibles
resumen_dotacion_filtrado["Cupos disponibles"] = (
    resumen_dotacion_filtrado[
        "Dotacion máxima autorizada"
    ]
    - resumen_dotacion_filtrado[
        "Vehículos actuales"
    ]
)

# Calcular porcentaje de ocupación por servicio
resumen_dotacion_filtrado["Porcentaje de ocupación"] = (
    resumen_dotacion_filtrado[
        "Vehículos actuales"
    ]
    .div(
        resumen_dotacion_filtrado[
            "Dotacion máxima autorizada"
        ].replace(0, pd.NA)
    )
    .mul(100)
    .round(1)
)

# Clasificar situación de la dotación
resumen_dotacion_filtrado["Situación"] = (
    resumen_dotacion_filtrado[
        "Cupos disponibles"
    ].apply(
        lambda valor: (
            "Sobre dotación"
            if valor < 0
            else "Dotación completa"
            if valor == 0
            else "Con cupos disponibles"
        )
    )
)

# Calcular indicadores generales
total_actual = resumen_dotacion_filtrado[
    "Vehículos actuales"
].sum()

total_autorizado = resumen_dotacion_filtrado[
    "Dotacion máxima autorizada"
].sum()

total_cupos = (
    total_autorizado
    - total_actual
)

porcentaje_ocupacion = (
    total_actual / total_autorizado * 100
    if total_autorizado > 0
    else 0
)

# Mostrar indicadores
dot1, dot2, dot3, dot4 = st.columns(4)

dot1.metric(
    "Vehículos actuales",
    f"{total_actual:,.0f}".replace(",", ".")
)

dot2.metric(
    "Dotación máxima autorizada",
    f"{total_autorizado:,.0f}".replace(",", ".")
)

dot3.metric(
    "Cupos disponibles",
    f"{total_cupos:,.0f}".replace(",", ".")
)

dot4.metric(
    "Ocupación",
    f"{porcentaje_ocupacion:.1f}%"
)

st.subheader("Indicadores principales")
col1, col2, col3, col4, col5 = st.columns(5)
col1.metric("Vehículos", f"{len(df_filtrado):,}".replace(",", "."))
col2.metric("Servicios", df_filtrado["Servicio"].nunique() if "Servicio" in df_filtrado.columns else 0)
altas = df_filtrado["Priorización"].astype(str).str.lower().str.contains("alta", na=False).sum() if "Priorización" in df_filtrado.columns else 0
col3.metric("Priorización alta", f"{altas:,}".replace(",", "."))
cuatro = df_filtrado["Cumple 4 criterios"].eq("cumple").sum() if "Cumple 4 criterios" in df_filtrado.columns else 0
col4.metric("4 criterios", f"{cuatro:,}".replace(",", "."))
gasto_total = df_filtrado.get("Gasto en mantención acumulada num", pd.Series(dtype=float)).sum(skipna=True)
col5.metric("Gasto mantención", formato_pesos(gasto_total))

st.subheader("Gráficos")
tab1, tab2, tab3, tab4, tab5 = st.tabs([
    "Vehículos por Servicio",
    "Mayor gasto",
    "Antigüedad flota",
    "4 criterios",
    "Antigüedad + alta",
])

with tab1:
    if "Servicio" in df_filtrado.columns:
        graf = df_filtrado.groupby("Servicio").size().reset_index(name="Cantidad").sort_values("Cantidad", ascending=False)
        st.plotly_chart(px.bar(graf, x="Servicio", y="Cantidad", title="Cantidad de vehículos por Servicio"), use_container_width=True)

with tab2:
    if "Gasto en mantención acumulada num" in df_filtrado.columns:
        cols = [c for c in ["Servicio", "I.R.N.V.M.", "Tipo vehículo", "Gasto en mantención acumulada num"] if c in df_filtrado.columns]
        top = df_filtrado[cols].dropna(subset=["Gasto en mantención acumulada num"]).sort_values("Gasto en mantención acumulada num", ascending=False).head(20)
        st.plotly_chart(px.bar(top, x="I.R.N.V.M.", y="Gasto en mantención acumulada num", color="Servicio", title="Top 20 vehículos con mayor gasto en mantención"), use_container_width=True)


with tab3:
    if "Antigüedad 2027" in df_filtrado.columns:
        antig = df_filtrado.dropna(
            subset=["Antigüedad 2027"]
        ).copy()

        antig = antig[
            antig["Antigüedad 2027"].between(0, 50)
        ]
        graf_antiguedad = px.histogram(
            antig,
            x="Antigüedad 2027",
            nbins=20,
            title="Distribución de antigüedad de la flota al año 2027",
            labels={
                "Antigüedad 2027": "Antigüedad del vehículo (años)",
            }
        )

        graf_antiguedad.update_xaxes(
            title_text="Antigüedad del vehículo (años)"
        )

        graf_antiguedad.update_yaxes(
            title_text="Cantidad"
        )
        graf_antiguedad.update_traces(
            hovertemplate=(
                "Antigüedad: %{x} años<br>"
                "Cantidad: %{y}"
                "<extra></extra>"
            )
        )

        st.plotly_chart(
            graf_antiguedad,
            use_container_width=True,
        )

with tab4:
    if "Cumple 4 criterios" in df_filtrado.columns:
        graf = df_filtrado.groupby("Cumple 4 criterios").size().reset_index(name="Cantidad")
        st.plotly_chart(px.bar(graf, x="Cumple 4 criterios", y="Cantidad", title="Vehículos que cumplen con 4 criterios"), use_container_width=True)

with tab5:
    graf = df_filtrado.groupby("Cumple antigüedad y priorización alta").size().reset_index(name="Cantidad")
    st.plotly_chart(px.bar(graf, x="Cumple antigüedad y priorización alta", y="Cantidad", title="Vehículos que cumplen antigüedad y priorización alta"), use_container_width=True)

st.subheader("Tabla de resultados")
df_mostrar = df_filtrado.drop(columns=[c for c in df_filtrado.columns if c.endswith(" num")], errors="ignore").copy()
if "Kilometraje acumulado" in df_mostrar.columns and "Kilometraje acumulado num" in df_filtrado.columns:
    df_mostrar["Kilometraje acumulado"] = df_filtrado["Kilometraje acumulado num"].apply(formato_miles)
if "Gasto en mantención acumulada" in df_mostrar.columns and "Gasto en mantención acumulada num" in df_filtrado.columns:
    df_mostrar["Gasto en mantención acumulada"] = df_filtrado["Gasto en mantención acumulada num"].apply(formato_pesos)


# Corrección para evitar error PyArrow por columnas con datos mixtos
columnas_texto = [
    "Servicio",
    "Año de Compra",
    "I.R.N.V.M.",
    "Año Vehículo",
    "Kilometraje acumulado",
    "Gasto en mantención acumulada",
]

for col in columnas_texto:
    if col in df_mostrar.columns:
        df_mostrar[col] = df_mostrar[col].astype(str)

# Mostrar Año Vehículo sin decimales
if "Año Vehículo" in df_mostrar.columns:
    df_mostrar["Año Vehículo"] = df_mostrar["Año Vehículo"].apply(
        lambda valor: (
            str(int(float(valor)))
            if pd.notna(valor)
            and str(valor).strip() not in ["", "No indica"]
            else "No indica"
        )
    )


# Evita errores PyArrow por columnas mixtas
df_mostrar = df_mostrar.fillna("").astype(str)

st.dataframe(df_mostrar, use_container_width=True, height=420)

st.subheader("Descargas")
col_a, col_b, col_c = st.columns(3)

with col_a:
    st.download_button(
        label="Descargar Excel filtrado",
        data=dataframe_a_excel(df_filtrado),
        file_name="reporte_catastro_filtrado.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

minuta = generar_minuta(
    df_filtrado,
    filtros_aplicados,
    resumen_dotacion_filtrado
)

with col_b:
    st.download_button(
        label="Descargar minuta TXT",
        data=minuta.encode("utf-8"),
        file_name="minuta_catastro_vehiculos.txt",
        mime="text/plain",
    )

with col_c:
    try:
        pdf_buffer = minuta_a_pdf_bytes(minuta)
        st.download_button(
            label="Descargar minuta PDF",
            data=pdf_buffer,
            file_name="minuta_catastro_vehiculos.pdf",
            mime="application/pdf",
        )
    except Exception as e:
        st.warning("No fue posible generar el PDF. Verifique que reportlab esté en requirements.txt.")

with st.expander("Ver minuta generada"):
    st.text(minuta)
